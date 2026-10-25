"""Build the editable progress report from its reviewed Markdown source."""
from pathlib import Path
import re
from datetime import datetime, timezone
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs/progress/2026-10-25'
SOURCE = OUT / 'Bao_cao_tien_do_25102026.md'


def inline(paragraph, text):
    for token in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        run = paragraph.add_run(token.strip('*') if token.startswith('**') else token.strip('`'))
        if token.startswith('**'):
            run.bold = True


def diagram():
    target = OUT / 'minh_chung/kien_truc_he_thong.png'
    image = Image.new('RGB', (1800, 510), 'white')
    d = ImageDraw.Draw(image)
    font = ImageFont.truetype('C:/Windows/Fonts/times.ttf', 31)
    bold = ImageFont.truetype('C:/Windows/Fonts/timesbd.ttf', 34)
    blocks = [(30, 'TRÌNH DUYỆT', 'React · TypeScript · Leaflet'),
              (640, 'REST API', 'Express · nghiệp vụ · phân quyền'),
              (1250, 'CƠ SỞ DỮ LIỆU', 'PostgreSQL · PostGIS')]
    for x, title, detail in blocks:
        d.rounded_rectangle((x, 35, x + 520, 175), radius=10, fill='#f2f3f5', outline='#333333', width=2)
        d.text((x + 260, 77), title, font=bold, fill='black', anchor='mm')
        d.text((x + 260, 128), detail, font=font, fill='black', anchor='mm')
    for x in (565, 1175):
        d.line((x, 105, x + 55, 105), fill='black', width=3)
        d.polygon(((x + 55, 105), (x + 43, 98), (x + 43, 112)), fill='black')
    d.text((295, 258), 'CAD GeoJSON và bản nháp thiết bị', font=font, fill='black', anchor='mm')
    d.line((295, 175, 295, 222), fill='#555555', width=2)
    d.text((900, 258), 'Supabase Auth và Storage khi cấu hình', font=font, fill='black', anchor='mm')
    d.line((900, 175, 900, 222), fill='#555555', width=2)
    d.text((900, 370), 'DWG → DXF/QGIS → GeoJSON → chuẩn hóa → tập tham chiếu WGS84', font=font, fill='black', anchor='mm')
    d.text((900, 432), 'Sơ đồ logic triển khai trong mã nguồn, không biểu diễn kết quả nghiệm thu môi trường thật', font=font, fill='#444444', anchor='mm')
    target.parent.mkdir(parents=True, exist_ok=True)
    image.save(target)


def build():
    diagram()
    doc = Document()
    sec = doc.sections[0]
    # Match the supplied outline: Letter portrait, Times New Roman, formal academic margins.
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(1.9)
    sec.footer_distance = Cm(0.9)
    sec.different_first_page_header_footer = True
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Times New Roman')
    normal = doc.styles['Normal']
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in [('Heading 1', 15), ('Heading 2', 12.5)]:
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(9)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    doc.styles['Title'].font.size = Pt(19)
    doc.styles['Title'].font.bold = True
    doc.styles['Caption'].font.size = Pt(10)
    doc.styles['Caption'].font.italic = True
    # The packaged template carries theme fonts and a blue Title border.
    # Remove both so native Word also uses the academic reference styling.
    for border in list(doc.styles.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    for fonts in doc.styles.element.xpath('.//w:rFonts'):
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():
                del fonts.attrib[attr]
        for slot in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            fonts.set(qn('w:' + slot), 'Times New Roman')
    footer = sec.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(Cm(16.8), WD_ALIGN_PARAGRAPH.RIGHT)
    footer.add_run('Tiến độ đến 25/10/2026\t')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for r in footer.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9)
    for text, size in [('TRƯỜNG ĐẠI HỌC TÀI NGUYÊN VÀ MÔI TRƯỜNG TP.HCM', 12),
                       ('KHOA CÔNG NGHỆ THÔNG TIN', 12)]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    p = doc.add_paragraph('NGUYỄN HOÀNG TUẤN')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(48)
    p.runs[0].bold = True; p.runs[0].font.size = Pt(15)
    p = doc.add_paragraph('BÁO CÁO TIẾN ĐỘ\nĐỒ ÁN TỐT NGHIỆP', 'Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p = doc.add_paragraph('XÂY DỰNG HỆ THỐNG WEBGIS HỖ TRỢ QUẢN LÝ\nHẠ TẦNG VÀ THEO DÕI SỰ CỐ CẤP NƯỚC')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(26)
    p.runs[0].font.size = Pt(15); p.runs[0].bold = True
    for text in ['Giảng viên hướng dẫn: ThS. Từ Thanh Trí',
                 'Sinh viên thực hiện: Nguyễn Hoàng Tuấn',
                 'Mã số sinh viên: 1150070049', 'Lớp: 11_TTMT',
                 'Kỳ báo cáo: 19/10/2026 đến 25/10/2026']:
        p = doc.add_paragraph(text); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph('TP. Hồ Chí Minh, ngày 25 tháng 10 năm 2026')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(25)
    doc.core_properties.title = 'Báo cáo tiến độ đồ án tốt nghiệp đến ngày 25 tháng 10 năm 2026'
    doc.core_properties.subject = 'WebGIS hỗ trợ quản lý hạ tầng và theo dõi sự cố cấp nước'
    doc.core_properties.author = 'Nguyễn Hoàng Tuấn'
    doc.core_properties.created = datetime(2026, 10, 25, tzinfo=timezone.utc)
    doc.core_properties.modified = datetime(2026, 10, 25, tzinfo=timezone.utc)
    chunks = SOURCE.read_text(encoding='utf-8').split('<!-- page -->')[1:]
    for chunk in chunks:
        first = True
        lines = chunk.strip().splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1; continue
            if line.startswith('## '):
                p = doc.add_paragraph(line[3:], 'Heading 1')
                if first: p.paragraph_format.page_break_before = True
                first = False
            elif line.startswith('### '):
                doc.add_paragraph(line[4:], 'Heading 2')
            elif line.startswith('!['):
                match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(OUT / match[2]), width=Cm(16.8))
                p = doc.add_paragraph(match[1], 'Caption'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif line.startswith('|'):
                rows = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    row = [x.strip() for x in lines[i].strip().strip('|').split('|')]
                    if not all(re.fullmatch(r':?-+:?', x) for x in row): rows.append(row)
                    i += 1
                table = doc.add_table(rows=0, cols=len(rows[0])); table.autofit = False
                table.style = 'Table Grid'
                widths = [Cm(5.3), Cm(11.7)] if len(rows[0]) == 2 else [Cm(17 / len(rows[0]))] * len(rows[0])
                for col, width in zip(table.columns, widths): col.width = width
                for ridx, values in enumerate(rows):
                    cells = table.add_row().cells
                    for cell, text, width in zip(cells, values, widths):
                        cell.width = width
                        tcpr = cell._tc.get_or_add_tcPr()
                        margins = OxmlElement('w:tcMar')
                        for edge in ['top', 'left', 'bottom', 'right']:
                            m = OxmlElement('w:' + edge); m.set(qn('w:w'), '85'); m.set(qn('w:type'), 'dxa'); margins.append(m)
                        tcpr.append(margins)
                        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.02
                        inline(p, text)
                        for run in p.runs: run.font.size = Pt(10.5)
                        if ridx == 0:
                            sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'EDEDED'); tcpr.append(sh)
                            for run in p.runs: run.bold = True
                    trpr = table.rows[-1]._tr.get_or_add_trPr()
                    trpr.append(OxmlElement('w:cantSplit'))
                    if ridx == 0: trpr.append(OxmlElement('w:tblHeader'))
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 0.5
                continue
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if 'https://' in line else WD_ALIGN_PARAGRAPH.JUSTIFY
                inline(p, line)
            i += 1
    target = OUT / 'Bao_cao_tien_do_25102026.docx'
    doc.save(target)
    print(target)


if __name__ == '__main__':
    build()
